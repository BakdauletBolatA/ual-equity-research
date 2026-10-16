import json,sys,warnings
warnings.filterwarnings('ignore')
from pycel import ExcelCompiler
import pycel.excellib as el, statistics
def _flat(args):
    out=[]
    for a in args:
        if hasattr(a,'flat'): out+= [x for x in a.flat]
        elif isinstance(a,(list,tuple)): out+=_flat(a)
        else: out.append(a)
    return out
def median(*args):
    v=[x for x in _flat(args) if isinstance(x,(int,float)) and not isinstance(x,bool)]
    return statistics.median(v)
def slope(ys,xs):
    y=[a for a in _flat([ys])]; x=[a for a in _flat([xs])]
    pairs=[(a,b) for a,b in zip(x,y) if isinstance(a,(int,float)) and isinstance(b,(int,float))]
    mx=sum(a for a,_ in pairs)/len(pairs); my=sum(b for _,b in pairs)/len(pairs)
    return sum((a-mx)*(b-my) for a,b in pairs)/sum((a-mx)**2 for a,_ in pairs)
el.median=median; el.slope=slope
from openpyxl import load_workbook
f=sys.argv[1]
xl=ExcelCompiler(filename=f)
wb=load_workbook(f)
errs=[];vals={}
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('='):
                addr=f"{ws.title}!{c.coordinate}"
                try:
                    v=xl.evaluate(addr)
                except Exception as e:
                    v=f"ERR {type(e).__name__}: {e}"[:150]
                if isinstance(v,str) and (v.startswith('ERR') or v.startswith('#')): errs.append((addr,c.value[:80],v))
                vals[addr]=v
reg=json.load(open(f+'.reg.json'))
out={}
for k,d in reg.items():
    out[k]={}
    for col,ref in d.items():
        if not isinstance(ref,str): continue
        sh,cell=ref.split('!'); sh=sh.strip("'"); cell=cell.replace('$','')
        a=f"{sh}!{cell}"
        if a in vals: out[k][col]=vals[a]
        else: out[k][col]=wb[sh][cell].value
json.dump(out,open(f+'.vals.json','w'),ensure_ascii=False,indent=0,default=str)
print('formulas',len(vals),'errors',len(errs))
for e in errs[:40]: print(e)
json.dump({k:v for k,v in vals.items()},open(f+'.all.json','w'),ensure_ascii=False,default=str)
