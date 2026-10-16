# -*- coding: utf-8 -*-
import json, sys, os
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

SCR, OUTDIR = sys.argv[1], sys.argv[2]
V = json.load(open(f"{SCR}/UAL_model.xlsx.vals.json"))
A = json.load(open(f"{SCR}/UAL_model.xlsx.all.json"))
CH = f"{SCR}/charts"; os.makedirs(CH, exist_ok=True)
def v(k, c=None):
    d = V[k]
    return d[c] if c is not None else list(d.values())[0]

# ---------- formatting helpers (one source for all docs)
def n0(x): return f"{x:,.0f}".replace(",", " ")
def n1(x): return f"{x:,.1f}".replace(",", " ").replace(".", ",")
def n2(x): return f"{x:,.2f}".replace(",", " ").replace(".", ",")
def pc(x, d=1): return (f"{x*100:.{d}f}%").replace(".", ",")
def spc(x, d=1): return ("+" if x >= 0 else "−") + pc(abs(x), d)
def usd(x, d=0): return ("−" if x < 0 else "") + "$" + (n0(abs(x)) if d == 0 else n2(abs(x)) if d == 2 else n1(abs(x)))
def bn(x): return "$" + n1(x / 1000) + " млрд"

PRICE = v("price", "2026E"); TGT = v("t_target"); LOW = v("t_low"); HIGH = v("t_high"); UP = v("t_up"); RATING = v("t_rating")
SV = V["sc_val"]; PROB = V["sc_prob"]
WACC = v("wacc"); KE = v("ke"); BETA = V["beta"]["v"]; G = v("g", "2026E")
MID = v("midcycle", "2025"); RVM = v("rv_impl_m"); IW = V["impl_wacc"]["v"]
EPS26 = v("Base_eps", "2026E"); EPS27 = v("Base_eps", "2027E")
TTM_EBITDAR = v("ebitdar", "TTM 2К26"); ND = v("net_debt", "1П26")
print("Target", TGT, LOW, HIGH, UP, RATING)

# ---------- chart style
TXT, TXT2, GRID = "#0b0b0b", "#52514e", "#e4e3df"
C1, C2, C3, C4 = "#2a78d6", "#eb6834", "#1baf7a", "#eda100"
plt.rcParams.update({"font.family": "Arial", "font.size": 10, "axes.edgecolor": GRID, "axes.labelcolor": TXT2,
                     "xtick.color": TXT2, "ytick.color": TXT2, "axes.spines.top": False, "axes.spines.right": False,
                     "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6, "axes.axisbelow": True})
def finish(fig, ax, src, name, ylabel=None, xlabel=None):
    if ylabel: ax.set_ylabel(ylabel)
    if xlabel: ax.set_xlabel(xlabel)
    ax.grid(axis="x", visible=False)
    fig.text(0.01, 0.01, "Источник: " + src, fontsize=7.5, color=TXT2)
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    p = f"{CH}/{name}.png"; fig.savefig(p, dpi=200); plt.close(fig); return p

YRS = ["2019", "2022", "2023", "2024", "2025", "TTM 2К26"]
YL = ["2019", "2022", "2023", "2024", "2025", "TTM\n2К26"]
charts = {}
# 1 revenue structure
fig, ax = plt.subplots(figsize=(8, 4.2))
cols = ["2019", "2020", "2022", "2023", "2024", "2025", "TTM 2К26"]
pax = [V["rev_pax"][c] / 1000 for c in cols]; car = [V["rev_cargo"][c] / 1000 for c in cols]; oth = [V["rev_oth"][c] / 1000 for c in cols]
x = range(len(cols))
ax.bar(x, pax, color=C1, width=0.6, label="Пассажирская", edgecolor="white", linewidth=1)
ax.bar(x, car, bottom=pax, color=C2, width=0.6, label="Грузовая", edgecolor="white", linewidth=1)
ax.bar(x, oth, bottom=[a + b for a, b in zip(pax, car)], color=C3, width=0.6, label="Прочая (в основном MileagePlus)", edgecolor="white", linewidth=1)
for i, c in enumerate(cols):
    ax.text(i, pax[i] + car[i] + oth[i] + 0.8, pc(V["mix_oth"][c]), ha="center", fontsize=8, color=TXT2)
ax.set_xticks(list(x), ["2019", "2020", "2022", "2023", "2024", "2025", "TTM\n2К26"])
ax.legend(frameon=False, fontsize=8, loc="upper left")
ax.set_ylim(0, 72)
charts["mix"] = finish(fig, ax, "10-K / 8-K UAL; подпись — доля прочей выручки. 2020: прочая −23% г/г против −70% у пассажирской.", "mix", ylabel="Выручка, $ млрд")
# 2 TRASM vs CASM-ex vs CASM
fig, ax = plt.subplots(figsize=(8, 4.2))
yc = ["2019", "2022", "2023", "2024", "2025", "1П25", "1П26"]
ax.plot(range(len(yc)), [V["trasm"][c] for c in yc], color=C1, lw=2, marker="o", ms=5, label="TRASM")
ax.plot(range(len(yc)), [V["casm"][c] for c in yc], color=C2, lw=2, marker="o", ms=5, label="CASM (вкл. топливо)")
ax.plot(range(len(yc)), [V["casmex"][c] for c in yc], color=C3, lw=2, marker="o", ms=5, label="CASM-ex")
for i, c in enumerate(yc):
    ax.text(i, V["trasm"][c] + 0.25, n2(V["trasm"][c]), ha="center", fontsize=7.5, color=TXT2)
    ax.text(i, V["casmex"][c] - 0.6, n2(V["casmex"][c]), ha="center", fontsize=7.5, color=TXT2)
ax.set_xticks(range(len(yc)), yc); ax.set_ylim(8, 21.5)
ax.legend(frameon=False, fontsize=8, loc="lower right", ncol=3)
charts["unit"] = finish(fig, ax, "8-K UAL по итогам 4 кв. 2019–2025 и 15.07.2026; CASM = (выручка − опер. прибыль)/ASM.", "unit", ylabel="центов на ASM")
# 3 quarterly margin and fuel: two small panels, separate axes (no dual axis)
QC = ["2К25", "3К25", "4К25", "1К26", "2К26"]
fig, (a1, a2) = plt.subplots(1, 2, figsize=(9, 4))
a1.bar(range(5), [V["q_fuel"][c] for c in QC], color=C2, width=0.55)
a1.plot(range(5), [V["q_gcjet"][c] for c in QC], color=TXT2, lw=1.5, ls="--", marker="o", ms=4, label="Спот Gulf Coast (EIA)")
for i, c in enumerate(QC): a1.text(i, V["q_fuel"][c] + 0.08, n2(V["q_fuel"][c]), ha="center", fontsize=8)
a1.set_title("Цена топлива UAL, $/галлон", fontsize=9, color=TXT, loc="left"); a1.set_xticks(range(5), QC); a1.legend(frameon=False, fontsize=7.5, loc="upper left"); a1.set_ylim(0, 4.9)
a1.grid(axis="x", visible=False)
a2.bar(range(5), [V["q_adjop"][c] * 100 for c in QC], color=C1, width=0.55)
for i, c in enumerate(QC): a2.text(i, V["q_adjop"][c] * 100 + 0.3, pc(V["q_adjop"][c]), ha="center", fontsize=8)
a2.set_title("Скорр. операционная маржа, %", fontsize=9, color=TXT, loc="left"); a2.set_xticks(range(5), QC); a2.set_ylim(0, 13.5)
charts["q"] = finish(fig, a2, "8-K UAL (квартальные релизы 2025–2026); FRED MJFUELUSGULF.", "q")
# 4 CFO NI FCF
fig, ax = plt.subplots(figsize=(8, 4.2))
yc = ["2019", "2022", "2023", "2024", "2025", "TTM 2К26"]
w_ = 0.26; xs = range(len(yc))
ax.bar([i - w_ for i in xs], [V["ni"][c] / 1000 for c in yc], w_, color=C1, label="Чистая прибыль", edgecolor="white")
ax.bar(list(xs), [V["cfo"][c] / 1000 for c in yc], w_, color=C3, label="CFO", edgecolor="white")
ax.bar([i + w_ for i in xs], [V["fcf"][c] / 1000 for c in yc], w_, color=C2, label="FCF = CFO − капзатраты", edgecolor="white")
for i, c in enumerate(yc):
    r = V["cfo_ni"][c]; ax.text(i, V["cfo"][c] / 1000 + 0.3, ("×" + n1(r)) if isinstance(r, float) else "", ha="center", fontsize=8, color=TXT2)
ax.axhline(0, color=TXT2, lw=0.8); ax.set_xticks(list(xs), YL); ax.legend(frameon=False, fontsize=8, loc="upper left")
charts["cf"] = finish(fig, ax, "XBRL 10-K/10-Q UAL; подпись — CFO / чистая прибыль.", "cf", ylabel="$ млрд")
# 5 leverage
fig, ax = plt.subplots(figsize=(8, 4.2))
yc = ["2019", "2022", "2023", "2024", "2025", "TTM 2К26"]
ax.plot(range(6), [V["nd_ebitdar"][c] for c in yc], color=C1, lw=2, marker="o", ms=6, label="Чистый скорр. долг / EBITDAR")
ax.plot(range(6), [V["nd_cfo"][c] for c in yc], color=C3, lw=2, marker="o", ms=6, label="Чистый скорр. долг / CFO")
for i, c in enumerate(yc):
    ax.text(i, V["nd_ebitdar"][c] + 0.12, n1(V["nd_ebitdar"][c]) + "x", ha="center", fontsize=8)
    ax.text(i, V["nd_cfo"][c] - 0.3, n1(V["nd_cfo"][c]) + "x", ha="center", fontsize=8, color=TXT2)
ax.set_xticks(range(6), YL); ax.set_ylim(0, 5); ax.legend(frameon=False, fontsize=8)
charts["lev"] = finish(fig, ax, "8-K UAL (Adjusted net debt); расчёт: долг + опер. аренда + пенсии − ден. средства. 2019 — оценка по XBRL.", "lev", ylabel="кратность, x")
# 6 football field
fig, ax = plt.subplots(figsize=(8.5, 4.4))
items = [
 ("DCF пессимистичный (25%)", v("Bear_psm", "2032E"), v("Bear_psg", "2032E")),
 ("DCF базовый (50%)", v("Base_psg", "2032E"), v("Base_psm", "2032E")),
 ("DCF оптимистичный (25%)", v("Bull_psm", "2032E"), v("Bull_psg", "2032E")),
 ("DCF база: WACC 7,5–9,5%, g 2,5%", A["Sensitivity!D9"], A["Sensitivity!D5"]),
 ("EV/EBITDAR 4,6–6,7x (история) × EBITDAR 2027E", None, None),
 ("Норм. EPS × P/E 4,1–15,0x (история)", None, None),
]
# multiples ranges
eb27 = v("Base_ebitdar", "2027E"); shs = v("shares", "2026E"); nol = v("nol", "2026E")
ev_lo = (v("hist_min", "EV/EBITDAR") * eb27 - ND + nol) / shs; ev_hi = (v("hist_max", "EV/EBITDAR") * eb27 - ND + nol) / shs
ne = v("im_normeps"); pe_lo = ne * v("hist_min", "P/E"); pe_hi = ne * v("hist_max", "P/E")
items[4] = (items[4][0], ev_lo, ev_hi); items[5] = (items[5][0], pe_lo, pe_hi)
for i, (lab, lo, hi) in enumerate(items[::-1]):
    lo2, hi2 = min(lo, hi), max(lo, hi)
    ax.barh(i, hi2 - lo2, left=lo2, color=C1 if "DCF" in lab else C3, height=0.5)
    ax.text(hi2 + 3, i, f"${lo2:.0f}–{hi2:.0f}", va="center", fontsize=8)
ax.set_yticks(range(len(items)), [x[0] for x in items[::-1]], fontsize=8)
ax.axvline(PRICE, color=TXT, lw=1.2, ls="--"); ax.text(PRICE + 1, len(items) - 0.4, f"Цена ${n2(PRICE)}", fontsize=8)
ax.axvline(TGT, color=C2, lw=1.6); ax.text(TGT + 1, -0.75, f"Цель ${TGT:.0f}", fontsize=8, color=TXT)
ax.set_xlim(0, 260); ax.grid(axis="y", visible=False); ax.grid(axis="x", visible=True)
fig.text(0.01, 0.01, "Источник: модель UAL_model.xlsx (листы DCF, Multiples, Sensitivity); границы DCF — методы Гордона и выходного мультипликатора.", fontsize=7.5, color=TXT2)
ax.set_xlabel("$ на акцию")
fig.tight_layout(rect=(0, 0.04, 1, 1)); p = f"{CH}/ff.png"; fig.savefig(p, dpi=200); plt.close(fig); charts["ff"] = p

# ---------- shared tables
peer_ops = [
 ["Показатель (2 кв. 2026)", "UAL", "DAL", "AAL", "LUV"],
 ["TRASM, центы (г/г)", "20,25 (+12,1%)", "22,45* (+12,4%)", "20,45 (+10,3%)", "17,91 (+16,2%)"],
 ["CASM-ex, центы (г/г)", "13,12 (+6,1%)", "14,09 (+6,8%)", "13,93 (+2,9%)", "12,45 (+3,4%)"],
 ["Топливо, $/галлон", "4,19", "3,93", "4,05", "3,92"],
 ["Скорр. опер. маржа", "5,4%", "8,8%", "2,7%", "6,7%"],
 ["Рост ASM г/г", "+3,5%", "~+1%", "+5,4%", "+0,2%"],
 ["Загрузка", "83,4%", "н/д", "83,2%", "79,3%"],
]
peer_ops_src = "8-K за 2 кв. 2026: UAL 15.07, DAL 10.07, AAL 23.07, LUV 22.07.2026. *DAL — скорр. TRASM без выручки НПЗ. LUV — RASM, короткие рейсы: несопоставим по удельным показателям."
targets = [
 ["Цель / прогноз менеджмента", "План", "Факт", "Статус"],
 ["Скорр. EPS 2023", "$10–12", "$10,05", "Выполнено (нижняя граница)"],
 ["Скорр. EPS 2024", "$9–11", "$10,61", "Выполнено"],
 ["Скорр. EPS 2025", "$11,50–13,50", "$10,62", "Не выполнено"],
 ["«Двузначная» маржа до налога (янв. 2025)", "≥10%", "7,8% (2025)", "Не выполнено; цель перенесена на 2027"],
 ["Скорр. EPS 2026 (янв. → апр. → июль)", "$12–14", "$7–11 → $9–11", "Снижен из-за топлива; консенсус $7,92"],
 ["Капзатраты 2025 / 2026", "<$7 / <$8 млрд", "$6,3 / ~$7,5 млрд", "Выполняется (помогают задержки Boeing)"],
 ["Инвестиционный рейтинг в 2026", "IG", "BB+ / Ba1 / BB+", "В процессе (S&P — позитивный прогноз)"],
]
targets_src = "Investor Updates UAL (8-K 21.01.2025, 20.01.2026, 21.04.2026, 15.07.2026); пресс-релизы по итогам 4 кв.; 10-Q 2К26."
fleet = [["Тип", "Контракт 2П26", "Ожид. 2П26", "Контракт 2027", "Ожид. 2027"]]
for k, lab in [("f787", "Boeing 787"), ("fmax9", "Boeing 737 MAX 9"), ("fmax10", "Boeing 737 MAX 10"), ("fa321", "Airbus A321neo"), ("fxlr", "Airbus A321XLR")]:
    d = V[k]; fleet.append([lab, n0(d["Контракт 2П26"]), n0(d["Ожид. 2П26"]), n0(d["Контракт 2027"]), n0(d["Ожид. 2027"])])
d = V["ftot"]; fleet.append(["Итого", n0(d["Контракт 2П26"]), n0(d["Ожид. 2П26"]), n0(d["Контракт 2027"]), n0(d["Ожид. 2027"])])
fleet_src = "10-Q UAL за 2 кв. 2026 (Firm Order and Option Aircraft). Ожидаемые поставки — с учётом задержек, сообщённых Boeing/Airbus."
def mfmt(x, f="x"):
    if not isinstance(x, (int, float)): return "н/п"
    return (n1(x) + "x") if f == "x" else pc(x) if f == "%" else n2(x) + "x"
mult = [["", "UAL", "DAL", "AAL", "LUV", "ALK", "JBLU", "Медиана"]]
for k, lab, f in [("m_evebitdar", "EV/EBITDAR (GAAP)", "x"), ("m_pe", "P/E TTM", "x"), ("m_fcfy", "FCF yield", "%"), ("m_evrev", "EV/выручка", "xx"), ("m_pb", "P/B", "xx"), ("m_opm", "Опер. маржа TTM", "%"), ("m_lev", "Чист. долг/EBITDAR", "x")]:
    mult.append([lab] + [mfmt(V[k].get(p), f) for p in ["UAL", "DAL", "AAL", "LUV", "ALK", "JBLU", "Медиана пиров"]])
mult.append(["P/E FY2026 (конс./гайденс)", mfmt(V["m_fpe"]["UAL"]), mfmt(V["m_fpe"]["DAL"]), "н/п", "н/д", "н/д", "н/д", "—"])
mult_src = "SEC XBRL (TTM на 30.06.2026: FY2025 + 1П26 − 1П25), цены 13.10.2026 (Yahoo Finance); EV включает опер. аренду и пенсии. Консенсус UAL — Yahoo Finance; DAL — гайденс 09.10.2026."
sens = [["Топливо 2027E, $/галл. \\ ΔTRASM", "−4%", "−2%", "0%", "+2%", "+4%"]]
for i, f_ in enumerate([2.70, 3.00, 3.30, 3.60, 3.90, 4.20]):
    sens.append([n2(f_)] + [usd(A[f"Sensitivity!{c}{20+i}"], 2) for c in "BCDEF"])
sens_src = "UAL_model.xlsx, лист Sensitivity (EPS 2027E, базовый сценарий: топливо $3,30, TRASM 19,68 цента; без переноса топлива в тариф)."
wg = [["WACC \\ g", "1,5%", "2,0%", "2,5%", "3,0%", "3,5%"]]
for i, w_ in enumerate([0.075, 0.08, 0.085, 0.09, 0.095]):
    wg.append([pc(w_)] + ["$" + n0(A[f"Sensitivity!{c}{5+i}"]) for c in "BCDEF"])
risks = [
 ["Риск", "Вероятн.", "Влияние", "Что отслеживать / порог"],
 ["Топливо и Ближний Восток (Brent $125 в начале окт.)", "Высокая", "Высокое", "Спот Gulf Coast >$4,50 устойчиво → пересмотр в Sell"],
 ["Рецессия США (UST10 5,2%, нефтяной шок)", "Средняя", "Высокое", "TRASM г/г ниже +5% в 4К26 / отрицательный в 2027"],
 ["Задержки Boeing (MAX 10 не сертифицирован)", "Высокая", "Среднее", "Поставки 787 в 2027 < 20; капзатраты сдвигаются, рост ASM ниже"],
 ["Труд: IAM, IBT открыты; пилоты — окт. 2027", "Средняя", "Среднее", "CASM-ex 2027 > +4% г/г"],
 ["FAA и инфраструктура (лимит ORD, ATC, Newark)", "Средняя", "Среднее", "Продление ограничений FAA на хабах"],
 ["Геополитика на международных маршрутах", "Средняя", "Среднее", "Доля МБВ/Индии/Африки уже −34% ASM в 2К26"],
 ["Долг и ставки (скорр. долг $34,7 млрд)", "Низкая", "Среднее", "Чистый долг / EBITDAR > 3,0x"],
 ["Конкуренция в премиуме (Delta, AA)", "Средняя", "Среднее", "Рост премиальной выручки UAL ниже DAL 2 кв. подряд"],
]
limits = [
 ["Слабое место", "Как влияет на вывод"],
 ["Терминальная стоимость > 100% EV в базовом DCF: FCFF отрицателен до 2029 из-за капзатрат 11–12% выручки", "±0,5 п.п. к росту TRASM в год меняют стоимость на ~±$50 → широкий диапазон; это аргумент за Hold, а не за уверенный Buy/Sell"],
 ["Путь цены топлива и доля переноса в тариф (2П26 $4,39; 2027 $3,30 по EIA)", "При $3,90 в 2027 EPS 2027E падает с $8,71 до $2,57; отчёт за 3К26 (20.10) может сдвинуть оценку"],
 ["Данные: 3К26 не вышел; пиры по GAAP со спецстатьями, НПЗ Delta, международные пиры не рассчитаны; доли хабов — вторичный источник", "Сравнение с пирами только вспомогательное; рекомендация опирается на DCF и историю UAL"],
]
facts3 = [
 f"Топливный шок съедает маржу: топливо $4,19/галл. во 2К26 (+79% г/г), скорр. опер. маржа 5,4% против 11,6% год назад; спот Gulf Coast в сент.–окт. $4,34–4,39 — выше допущения UAL на 3К ($3,69). Наша EPS 2026E ${n2(EPS26)} против консенсуса $7,92.",
 f"Франшиза выдерживает шок: TRASM +12,1% во 2К26, премиум +16%, MileagePlus ~11% выручки; чистый долг/EBITDAR {n1(v('nd_ebitdar','TTM 2К26'))}x, ликвидность $19,6 млрд.",
 f"Цена уже закладывает возврат к средней марже цикла: reverse DCF даёт маржу 2032E {pc(RVM)} при средней за цикл {pc(MID)}; для роста нужна маржа выше истории, а цели менеджмента в 2025–2026 не выполнены.",
]
monitor = [
 ["Показатель", "Сейчас", "Порог пересмотра"],
 ["Спот jet fuel Gulf Coast", "$4,34 (06.10.2026)", "<$3,00 → Buy; >$4,50 → Sell"],
 ["TRASM г/г (4К26)", "+12,1% (2К26)", "<+10% → Sell (тариф не покрывает топливо)"],
 ["CASM-ex г/г (прогноз 2027)", "+6,0% (1П26)", ">+4% → вниз; <+2% → вверх"],
 ["Скорр. маржа до налога 2027 (прогноз UAL)", "цель ≥10%", "Подтверждение ≥10% при топливе <$3,20 → Buy"],
 ["Чистый долг / EBITDAR", f"{n1(v('nd_ebitdar','TTM 2К26'))}x", ">3,0x → Sell; рейтинг IG → вверх"],
]
json.dump(dict(target=TGT, low=LOW, high=HIGH, up=UP, rating=RATING), open(f"{SCR}/summary.json", "w"))

# ================================================================= PPTX
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
NAVY = RGBColor(0x1F, 0x38, 0x64); GREY = RGBColor(0x52, 0x51, 0x4E); BLK = RGBColor(0x0B, 0x0B, 0x0B); LIGHT = RGBColor(0xD9, 0xE1, 0xF2)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BL = prs.slide_layouts[6]
def tb(sl, x, y, w, h, text, size=12, bold=False, color=BLK, align=None):
    t = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color; p.font.name = "Arial"
        if align: p.alignment = align
        p.space_after = Pt(4)
    return t
def slide(title, num, src=None):
    sl = prs.slides.add_slide(BL)
    bar = sl.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12)); bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb(sl, 0.5, 0.3, 12.3, 1.0, title, 22, True, NAVY)
    tb(sl, 0.5, 7.0, 10.5, 0.4, ("Источник: " + src) if src else "", 9, False, GREY)
    tb(sl, 11.3, 7.0, 1.6, 0.4, f"UAL | {num}", 9, False, GREY, PP_ALIGN.RIGHT)
    return sl
def table(sl, rows, x, y, w, h, fs=11, colw=None):
    t = sl.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for i, r in enumerate(rows):
        for j, val in enumerate(r):
            c = t.cell(i, j); c.text = str(val)
            p = c.text_frame.paragraphs[0]; p.font.size = Pt(fs); p.font.name = "Arial"
            p.font.bold = (i == 0); p.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else BLK
            c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 == 0 else RGBColor(255, 255, 255))
    if colw:
        for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    return t
def pic(sl, p, x=0.9, y=1.35, w=11.5):
    sl.shapes.add_picture(p, Inches(x), Inches(y), width=Inches(w))

# 1 recommendation
sl = slide(f"{RATING}: цель ${TGT:.0f} (+{UP*100:.0f}%) — франшиза сильная, но цена уже учитывает возврат маржи к средней за цикл", 1,
           "UAL_model.xlsx; цена NASDAQ на закрытие 13.10.2026; отчётность по 2 кв. 2026 (3 кв. выходит 20.10.2026).")
box = sl.shapes.add_shape(1, Inches(0.5), Inches(1.5), Inches(4.2), Inches(5.2)); box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.fill.background()
tb(sl, 0.7, 1.65, 3.9, 5.0, [
    "United Airlines (NASDAQ: UAL)", f"Рекомендация: {RATING}", f"Цель 12 мес.: ${TGT:.0f}", f"Диапазон: ${LOW:.0f}–{HIGH:.0f}",
    f"Цена: ${n2(PRICE)} (13.10.2026)", f"Потенциал: {spc(UP)}", "",
    f"Сценарии: ${SV['Bear']:.0f} / ${SV['Base']:.0f} / ${SV['Bull']:.0f}", "Вероятности: 25% / 50% / 25%",
    f"WACC {pc(WACC)}, g {pc(G)}", f"EPS 2026E (наш): ${n2(EPS26)}", f"EPS 2027E (наш): ${n2(EPS27)}"], 14, False, BLK)
tb(sl, 5.1, 1.55, 7.8, 0.5, "Три факта, на которых держится вывод", 15, True, NAVY)
tb(sl, 5.1, 2.1, 7.8, 4.6, [f"{i+1}. {t}" for i, t in enumerate(facts3)], 13)
# 2 business
sl = slide("Прочая выручка (в основном MileagePlus) — 6,5% выручки, но в 2020 она упала лишь на 23% против −70% у билетов: это стабилизатор", 2,
           "10-K 2025 (прим. 2: партнёрская выручка $3,2 млрд, погашения миль $3,4 млрд); 8-K UAL.")
pic(sl, charts["mix"], 0.5, 1.4, 8.0)
tb(sl, 8.7, 1.5, 4.3, 5.4, [
    "Хабы: Newark, Chicago, Denver, Houston, San Francisco, Washington-Dulles, Los Angeles, Guam; Star Alliance + JV с Lufthansa/Air Canada, ANA, Air NZ.",
    "Доля пассажиров в хабе (BTS 2024, вторичн. источник): EWR 58%, SFO 44%, IAH 39%, DEN 38%, ORD 30%.",
    f"MileagePlus: партнёрская выручка $3,2 млрд + погашения на перелёты $3,4 млрд ≈ {pc(v('loy_total_pct','2025'))} выручки 2025 (наша оценка).",
    "Chase платит за мили вперёд по объёму трат по картам — выручка зависит от потребительских расходов, а не от продаж кресел; договор продлён до 2029.",
    "Груз: 3,0% выручки, +23% г/г во 2К26.", "Слабость: доля премиальной выручки UAL не раскрывается (не найдено) — у Delta премиум ≈39% выручки."], 11)
# 3 targets
sl = slide("Менеджмент выполняет прогнозы капзатрат, но не прибыли: цели по EPS 2025 и марже 2025 пропущены, прогноз 2026 снижен", 3, targets_src)
table(sl, targets, 0.6, 1.5, 12.1, 4.6, 12, [4.2, 2.0, 2.4, 3.5])
tb(sl, 0.6, 6.25, 12, 0.7, "Стратегия United Next: премиальные кабины (Polaris, Premium Plus), более 630 новых самолётов до 2034, Starlink во всём парке к концу 2027 (450 бортов на июль 2026), A321XLR с международными рейсами с начала 2027.", 11, False, GREY)
# 4 peers ops
sl = slide("Рост RASM у UAL на уровне Delta, но маржа ниже: дорогое топливо ($4,19/галл.) и CASM-ex +6% — преимущество над DAL не доказано", 4, peer_ops_src)
table(sl, peer_ops, 0.8, 1.5, 11.7, 4.3, 13, [3.3, 2.1, 2.1, 2.1, 2.1])
tb(sl, 0.8, 6.0, 11.7, 0.9, "Пунктуальность: точные данные DOT по перевозчикам за 2025 не найдены; PIRG (по данным DOT) относит к лучшим Delta и Southwest, United в тройку не входит. UAL сообщает о лучших вылетах вовремя за 2 кв. с 2021 года.", 11, False, GREY)
# 5 unit economics
sl = slide("С 2019 CASM-ex вырос на 24%, а TRASM — на 18%: затраты обгоняют выручку, поэтому маржа 2025 ниже докризисной", 5,
           "8-K UAL по итогам 4 кв. 2019–2025 и 15.07.2026.")
pic(sl, charts["unit"], 0.5, 1.4, 8.0)
tb(sl, 8.7, 1.5, 4.3, 5.4, [
    f"Скорр. опер. маржа: 2019 {pc(V['adj_opm']['2019'])} → 2023 {pc(V['adj_opm']['2023'])} → 2025 {pc(V['adj_opm']['2025'])} → TTM {pc(V['adj_opm']['TTM 2К26'])}.",
    f"Средняя за цикл (2019, 2022–2025): {pc(MID)}.",
    f"Загрузка 2025: {pc(V['lf']['2025'])}; доходность {n2(V['yield']['2025'])} цента на RPM.",
    f"Топливная эффективность: {n2(V['gpa']['2019'])} → {n2(V['gpa']['2025'])} галл. на 1000 ASM (−1,1% в год, новые самолёты).",
    f"ASM на сотрудника: {n2(V['asm_emp']['2019'])} млн (2019) → {n2(V['asm_emp']['2025'])} млн (2025) — производительность вернулась к уровню 2019."], 11)
# 6 fuel
sl = slide("Скачок цены топлива в 2026 съел половину маржи: тариф покрывает рост с лагом (2К26 ~50%), а спот в сент.–окт. выше прогноза UAL", 6,
           "8-K UAL; FRED MJFUELUSGULF / DJFUELUSGULF (EIA); Investor Update 15.07.2026.")
pic(sl, charts["q"], 0.5, 1.35, 8.4)
tb(sl, 9.1, 1.5, 3.9, 5.4, [
    "UAL не хеджирует топливо: +10% к цене = −$1,4 млрд опер. прибыли за TTM (33% скорр. опер. прибыли); ±1% RASM = ±$0,63 млрд (15%).",
    "Прогноз UAL на 3К26 исходил из $3,69/галл.; спот Gulf Coast в сентябре $4,39 (+$0,35 логистика и налоги) → риск недовыполнения.",
    "Delta 09.10.2026 снизила прогноз EPS на 2026 до $5,10–5,60 (с $6,50–7,50) из-за топлива — сигнал для всего сектора."], 11)
# 7 cash flow
sl = slide(f"CFO в {n1(V['cfo_ni']['2025'])} раза выше прибыли из-за амортизации, неденежного налога (NOL) и авансов; FCF съедают капзатраты ~11% выручки", 7,
           "XBRL 10-K/10-Q UAL; балансы 8-K (ATL и отложенная выручка MileagePlus).")
pic(sl, charts["cf"], 0.5, 1.4, 8.0)
tb(sl, 8.7, 1.5, 4.3, 5.4, [
    f"2025: прирост авансов (ATL) и отложенной выручки миль дал {bn(V['d_atl']['2025'])} к CFO; без него CFO/прибыль = {n1(V['cfoex_ni']['2025'])}x.",
    f"(ATL + отложенная выручка) = {pc(V['liab_rev']['2025'])} выручки: при спаде спроса этот приток разворачивается.",
    "Текущий налог 2025 — $28 млн при прибыли до налога $4,3 млрд: NOL $10,6 млрд.",
    f"Скорр. капзатраты: {pc(V['capex_rev']['2025'])} выручки в 2025, ~$7,5 млрд в 2026.",
    f"Выкуп акций: $637 млн в 2025, $27 млн в 1П26; остаток программы $755 млн."], 11)
# 8 debt
sl = slide(f"Долг под контролем: чистый скорр. долг ${n1(ND/1000)} млрд = {n1(v('nd_ebitdar','TTM 2К26'))}x EBITDAR, ликвидность $19,6 млрд; нагрузка через CFO выглядит мягче, чем есть", 8,
           "8-K 15.07.2026; 10-K 2025 (график обязательств); 10-Q 2К26 (рейтинги).")
pic(sl, charts["lev"], 0.5, 1.4, 8.0)
tb(sl, 8.7, 1.5, 4.3, 5.4, [
    f"Скорр. долг ${n1(v('adj_debt','1П26')/1000)} млрд: долг ${n1(V['debt']['1П26']/1000)} + аренда ${n1(V['oplease']['1П26']/1000)} + пенсии ${n1(V['pension']['1П26']/1000)}.",
    f"Платежи по долгу и аренде: ${n1(V['mat_tot']['2026'])} млрд в 2026, ${n1(V['mat_tot']['2027'])} млрд в 2027.",
    f"Покрытие процентов (скорр. EBIT): {n1(V['cov']['TTM 2К26'])}x; EBITDAR/(% + аренда): {n1(V['covr']['TTM 2К26'])}x.",
    "Рейтинги BB+ / Ba1 / BB+, цель — IG в 2026.",
    f"Долг/CFO ({n1(V['nd_cfo']['TTM 2К26'])}x) занижен авансами и нулевыми налогами; надёжнее EBITDAR через цикл (без ATL CFO даёт {n1(V['nd_cfoex']['2025'])}x в 2025)."], 11)
# 9 fleet
sl = slide(f"Boeing поставит UAL лишь ~{pc(v('fboe'),0)} контрактных самолётов в 2П26–2027: рост ограничен, но капзатраты ниже плана", 9, fleet_src)
table(sl, fleet, 0.8, 1.5, 8.0, 3.8, 13, [2.8, 1.3, 1.3, 1.3, 1.3])
tb(sl, 9.1, 1.5, 3.9, 5.3, [
    "MAX 10 не сертифицирован: 0 поставок в 2П26 вместо 3; в 2027 ожидается 20 из 44.",
    "787: 12 ожидаемых в 2П26 против 42 по контракту.",
    "Последствия: рост ASM 2026 ~2,5%; капзатраты сдвигаются на 2028+ (по контракту $7,4–9,0 млрд в год в 2028–2030).",
    "Плюс: меньше капзатрат и выше загрузка в 2026–2027. Минус: старый парк дольше летает, CASM-ex выше."], 11)
# 10 multiples
sl = slide("UAL дешевле пиров по EV/EBITDAR (6,3x против 9,0x), но медиана искажена убытками; против своей истории дешевизны нет", 10, mult_src)
table(sl, mult, 0.5, 1.45, 12.3, 4.4, 11, [2.9, 1.3, 1.3, 1.3, 1.3, 1.3, 1.3, 1.6])
tb(sl, 0.5, 6.0, 12.3, 0.9, f"История UAL (EV/EBITDAR на конец года, 2019 и 2022–2025): {n1(v('hist_min','EV/EBITDAR'))}–{n1(v('hist_max','EV/EBITDAR'))}x, медиана {n1(v('hist_med','EV/EBITDAR'))}x; сейчас {n1(V['hm_evebitdar']['Текущий'])}x (то же определение) — верх диапазона. P/E на нормализованной прибыли — {n1(v('im_pe_norm'))}x.", 11, False, GREY)
# 11 valuation
sl = slide(f"Взвешенная DCF-оценка ${SV['Взвешенно']:.0f} лишь на {pc(SV['Взвешенно']/PRICE-1,0)} выше цены; разброс сценариев ${SV['Bear']:.0f}–{SV['Bull']:.0f} отражает высокий операционный рычаг", 11,
           "UAL_model.xlsx (DCF, Multiples, Sensitivity, ReverseDCF).")
pic(sl, charts["ff"], 0.4, 1.35, 8.3)
tb(sl, 8.9, 1.45, 4.1, 5.4, [
    f"WACC {pc(WACC)}: Rf 5,24%, β {n2(BETA)}, ERP 5%, Kd 6,5% до налога; долг с арендой {pc(v('w_wd'),0)} капитала.",
    f"Терминал: g {pc(G)}, RONIC 10%; выход по EV/EBITDAR {n1(v('hist_med','EV/EBITDAR'))}x (база).",
    f"Reverse DCF: цена закладывает маржу 2032E {pc(RVM)} ≈ средней за цикл {pc(MID)} и WACC {pc(IW)}.",
    f"Наша EPS 2027E ${n2(EPS27)} против консенсуса $12,83: консенсус ждёт возврата к средней марже уже в 2027, мы — к 2030–2032."], 11)
# 12 sensitivity
sl = slide("Топливо — главный драйвер: каждые $0,30/галл. в 2027 стоят ~$2,9 EPS, а 2% TRASM — ~$2,9; при $3,90 и неизменном RASM EPS 2027E падает до $2,6", 12, sens_src)
table(sl, sens, 0.8, 1.5, 7.6, 4.2, 13, [2.6, 1.0, 1.0, 1.0, 1.0, 1.0])
tb(sl, 8.8, 1.5, 4.2, 5.0, ["Как читать: строка — цена топлива 2027E, столбец — отклонение TRASM 2027E от базы; эффект статический (без переноса в тариф).",
    f"Стоимость акции (Гордон) по WACC 7,5–9,5% при g 2,5%: ${A['Sensitivity!D9']:.0f}–{A['Sensitivity!D5']:.0f}; ±0,5 п.п. WACC ≈ ±$15–20. Полная таблица WACC × g — в отчёте и Excel (лист Sensitivity).",
    "Перенос топлива в тариф — ключ: UAL заявляет 100% к 4К26; Delta и Deutsche Bank ждут полного переноса только в 2027."], 11)
# 13 risks
sl = slide("Главные риски — топливо и спрос; вывод неверен при топливе ниже $3/галл. (апсайд к ~$200) или росте TRASM 4К26 менее 10%", 13,
           "Оценка аналитика; 10-K 2025 (Risk Factors), 10-Q 2К26.")
table(sl, risks, 0.5, 1.4, 12.3, 4.4, 11, [4.6, 1.2, 1.2, 5.3])
tb(sl, 0.5, 5.95, 12.3, 1.0, "Катализаторы: отчёт за 3К26 и прогноз на 4К (20.10.2026); снижение нефти (EIA: Brent $84 в 2027); инвестиционный рейтинг (2026); A321XLR на международных рейсах (начало 2027); возобновление выкупа ($755 млн); сертификация MAX 10.", 11, False, GREY)
# 14 limitations
sl = slide("Ограничения: оценка держится на терминальной марже и цене топлива, поэтому вывод — Hold с широким диапазоном", 14, "Оценка аналитика.")
table(sl, limits, 0.5, 1.5, 12.3, 4.2, 13, [5.6, 6.7])
tb(sl, 0.5, 6.0, 12.3, 0.9, "Прочее: бета 1,36 сырая (по Блюму 1,24 → WACC ~8,2%, оценка +~$10); консенсус EPS — с агрегатора; скорр. долг 2019 — оценка по XBRL.", 11, False, GREY)
prs.save(f"{OUTDIR}/UAL_presentation.pptx")
print("pptx ok")

# ================================================================= DOCX
from docx import Document
from docx.shared import Pt as DPt, Cm, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc = Document()
sec = doc.sections[0]; sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = DPt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
for hs in ["Heading 1", "Heading 2"]:
    s = doc.styles[hs]; s.font.name = "Arial"; s.font.color.rgb = DRGB(0x1F, 0x38, 0x64); s.font.size = DPt(14 if hs == "Heading 1" else 11.5)
def P(text, bold=False, size=None, italic=False, color=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = DPt(size)
    if color: r.font.color.rgb = DRGB(*color)
    p.paragraph_format.space_after = DPt(4); return p
def B(text):
    p = doc.add_paragraph(text, style="List Bullet"); p.paragraph_format.space_after = DPt(2); return p
def SRC(text): P("Источник: " + text, italic=True, size=8, color=(0x52, 0x51, 0x4E))
def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)
def T(rows, widths=None, fs=8.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, val in enumerate(r):
            c = t.cell(i, j); c.text = ""; run = c.paragraphs[0].add_run(str(val)); run.font.size = DPt(fs); run.bold = (i == 0)
            if i == 0: shade(c, "D9E1F2")
    if widths:
        for j, w in enumerate(widths):
            for r in t.rows: r.cells[j].width = Cm(w)
    return t
def IMG(p, w=16): doc.add_picture(p, width=Cm(w))

doc.add_heading("United Airlines Holdings (NASDAQ: UAL): " + RATING + ", целевая цена $" + f"{TGT:.0f}", 0)
P(f"Дата анализа: 14.10.2026. Цена: ${n2(PRICE)} (закрытие 13.10.2026). Последняя отчётность: 10-Q за 2 кв. 2026 (16.07.2026). Отчёт за 3 кв. 2026 выйдет 20.10.2026 и в анализе не используется. Валюта — доллары США; финансовые данные в млн долл., если не указано иное.", italic=True, size=9)
P("Разметка: [Ф] — факт из отчётности, [Д] — допущение, [О] — оценка аналитика. Все расчёты — в UAL_model.xlsx; номера листов указаны в тексте.", italic=True, size=9)
doc.add_heading("Резюме", 1)
T([["Рекомендация", "Цель 12 мес.", "Диапазон", "Потенциал", "Сценарии (25/50/25%)"],
   [RATING, f"${TGT:.0f}", f"${LOW:.0f}–{HIGH:.0f}", spc(UP), f"${SV['Bear']:.0f} / ${SV['Base']:.0f} / ${SV['Bull']:.0f}"]], [3, 3, 3, 3, 5], 9.5)
P("")
P("Вывод держится на трёх фактах:", bold=True)
for t in facts3: B(t)
P(f"Почему не Buy: в цене $103 уже заложен возврат маржи к средней за цикл ({pc(MID)}), а для роста нужна маржа выше истории при топливе дороже $4. Почему не Sell: баланс и ликвидность выдерживают шок, выручка перекладывает топливо в тариф (TRASM +12%), а сценарий нормализации нефти (EIA: Brent $84 в 2027) даёт оптимистичную оценку около $200. По правилу рейтинга (Buy при потенциале > +15%, Sell при < −10%) потенциал {spc(UP)} соответствует Hold.")

doc.add_heading("1. Бизнес и стратегия", 1)
doc.add_heading("Модель", 2)
P(f"[Ф] United — крупнейшая авиакомпания мира по ASM: {n0(V['asm']['2025'])} млн ASM в 2025, {n0(V['fleet']['1П26'])} самолётов на 30.06.2026 (1 122 магистральных и 430 региональных), {n1(V['emp']['1П26'])} тыс. сотрудников. Хабы — Chicago O'Hare, Denver, Houston, Los Angeles, Newark, San Francisco, Washington-Dulles и Guam (10-K 2025). United входит в Star Alliance и в совместные предприятия с Air Canada и Lufthansa Group (Атлантика), ANA (Тихий океан) и Air New Zealand. Международные перевозки — 45% ASM во 2К26 (38 949 из 87 279 млн): Атлантика 18 579, Тихий океан 11 487, Латинская Америка 8 883 млн ASM (8-K 15.07.2026).")
P("[Ф, вторичный источник] Доля United в пассажиропотоке хабов (BTS за 2024, по данным Simple Flying): Newark 58%, San Francisco 44%, Houston 39%, Denver 38%, Chicago 30%. Washington-Dulles — не найдено. Доминирование в Newark и SFO — главный актив, но в Chicago и Denver United делит хаб с American и Southwest, там преимущество слабее. Слоты: Newark, JFK, LGA и DCA регулируются, плюс в 2025–2026 действует ограничение FAA на O'Hare (Investor Update 15.07.2026).")
IMG(charts["mix"], 15.5); SRC("10-K и 8-K UAL; лист Hist / Metrics.")
doc.add_heading("MileagePlus и Chase", 2)
P(f"[Ф] В 2025 United признал в прочей выручке $3,2 млрд от партнёров MileagePlus (маркетинг, реклама, неавиационные погашения), против $2,9 млрд в 2024 и $2,7 млрд в 2023. Ещё $3,4 млрд погашений миль на перелёты учтены в пассажирской выручке, начислено миль на $3,9 млрд (10-K 2025, прим. 2). [О] В сумме MileagePlus — около ${n1(v('loy_total','2025')/1000)} млрд, или {pc(v('loy_total_pct','2025'))} выручки. Это нижняя оценка: прибыль программы не раскрывается — не найдено. Ключевой партнёр — JPMorgan Chase, договор продлён в 2020 году до 2029 (вторичный источник: The Points Guy). Выручка лояльности выросла на 9% в 2025 и на 11% во 2К26.")
P("Почему выручка устойчивее: Chase платит за мили при тратах по картам, то есть по потребительским расходам, а не по проданным креслам. Деньги приходят вперёд, а выручка признаётся позже — отложенная выручка MileagePlus составляет $7,8 млрд. Проверка на стресс: в 2020 пассажирская выручка упала на 70%, а прочая — только на 23% ($2 455 → $1 902 млн, лист Hist). Риск — регулирование межбанковских комиссий по картам (10-K, Risk Factors).")
doc.add_heading("Стратегия и выполнение целей", 2)
P("United Next — это премиальные кабины (Polaris, Polaris Studio, Premium Plus), более 630 новых самолётов до 2034, рост на дальних маршрутах (A321XLR в Европу и Южную Америку с начала 2027) и цифровые продукты (Starlink бесплатно для участников MileagePlus: 450 бортов на июль 2026, весь парк — к концу 2027; медиасеть Kinective). [Ф] Во 2К26 премиальная выручка выросла на 16%, корпоративная — на 27%, Basic Economy — на 11% (8-K 15.07.2026). Долю премиума в выручке UAL не раскрывает — не найдено.")
T(targets, [5.2, 2.6, 3.0, 5.2]); SRC(targets_src)
P("Вывод: прогнозы капзатрат выполняются, а прибыльные цели в 2025–2026 — нет. Отчасти это внешние шоки (тарифы и спрос в 2025, нефть в 2026), но цель по двузначной марже переносится уже второй год. Поэтому прогнозам менеджмента мы доверяем с дисконтом и в базовом сценарии не закладываем маржу ≥10% в 2027.")
doc.add_heading("Конкурентные преимущества — доказательства и слабые места", 2)
T(peer_ops, [4.4, 3.0, 3.0, 3.0, 3.0]); SRC(peer_ops_src)
P("Сильные стороны: самая большая международная сеть среди американских перевозчиков, доминирование в Newark и SFO, рост RASM на уровне Delta (+12% во 2К26), лояльность и груз как стабилизаторы. Слабые: маржа ниже Delta (5,4% против 8,8%) из-за более дорогого топлива (у Delta свой НПЗ) и роста CASM-ex на 6%; пунктуальность за 2025 по данным DOT не найдена, а по рейтингу PIRG United не входит в тройку лучших.")

doc.add_heading("2. Операционные и финансовые показатели", 1)
yy = ["2019", "2022", "2023", "2024", "2025", "1П26", "TTM 2К26"]
def rowv(lab, key, f):
    return [lab] + [f(V[key][c]) if c in V[key] and isinstance(V[key][c], (int, float)) else "н/д" for c in yy]
ops = [["Показатель"] + yy,
       rowv("ASM, млрд", "asm", lambda x: n1(x / 1000)), rowv("RPM, млрд", "rpm", lambda x: n1(x / 1000)), rowv("Загрузка", "lf", pc),
       rowv("Yield, центы", "yield", n2), rowv("PRASM, центы", "prasm", n2), rowv("TRASM, центы", "trasm", n2), rowv("CASM, центы", "casm", n2),
       rowv("CASM-ex, центы", "casmex", n2), rowv("Топливо, $/галл.", "fuelpx", n2), rowv("Галл./1000 ASM", "gpa", n2),
       rowv("Выручка, $ млрд", "rev", lambda x: n1(x / 1000)), rowv("Скорр. опер. маржа", "adj_opm", pc), rowv("Чистая прибыль, $ млрд", "ni", lambda x: n2(x / 1000)),
       rowv("EBITDAR, $ млрд", "ebitdar", lambda x: n2(x / 1000)), rowv("FCF (CFO−capex), $ млрд", "fcf", lambda x: n2(x / 1000)), rowv("Capex скорр./выручка", "capex_rev", pc)]
T(ops, [3.6] + [1.9] * 7, 8); SRC("листы Hist и Metrics (8-K и XBRL UAL). 1П26 — полугодие; TTM = 2025 + 1П26 − 1П25. Топливо и CASM-ex за TTM не суммируются (н/д).")
P(f"Динамика. Выручка 2025 — $59,1 млрд (+37% к 2019), но скорр. опер. маржа {pc(V['adj_opm']['2025'])} ниже уровня 2019 ({pc(V['adj_opm']['2019'])}). CASM-ex вырос на 24% (10,21 → 12,64 цента), TRASM — на 18% (15,18 → 17,88). Основные причины — новые коллективные договоры (зарплаты +46% при ASM +16%) и аэропортовые сборы. В 2026 топливо подорожало до $3,53 за 1П (во 2К — $4,19), TRASM выросла на 9,7% за 1П, CASM-ex — на 6,0%.")
IMG(charts["unit"], 15); SRC("8-K UAL; лист Metrics.")
IMG(charts["q"], 15.5); SRC("8-K UAL; FRED (EIA).")
doc.add_heading("Капзатраты, поставки и Boeing", 2)
T(fleet, [4.0, 2.6, 2.6, 2.6, 2.6]); SRC(fleet_src)
P(f"[О] Из контрактных поставок Boeing на 2П26–2027 UAL ожидает лишь {pc(v('fboe'),0)}. MAX 10 не сертифицирован, 787 сдвинуты. Эффект двоякий: рост мощностей ниже (ASM 2026 ~+2,5%) и капзатраты 2026 ниже — ~$7,5 млрд против «<$8 млрд» в январе. Но старый парк (A319/A320, 757/767) летает дольше, топливная эффективность растёт медленнее, а обязательства переезжают в 2028–2030 ($7,4–9,0 млрд в год по контрактам, 10-K 2025). В модели капзатраты 2027–2030 — 11–12% выручки против 10,6% в 2025.")
doc.add_heading("Качество прибыли", 2)
IMG(charts["cf"], 15); SRC("XBRL 10-K/10-Q UAL.")
P(f"CFO/прибыль: {n1(V['cfo_ni']['2025'])}x в 2025. Разница с прибылью — это амортизация ($2,9 млрд), отложенный налог ($0,9 млрд неденежного налога: денежный налог $28 млн благодаря NOL) и прирост авансов и отложенной выручки миль ({bn(V['d_atl']['2025'])}). Без последнего CFO/прибыль = {n1(V['cfoex_ni']['2025'])}x. Авансы равны {pc(V['liab_rev']['2025'])} выручки и при спаде продаж уходят обратно: в 2023 их прирост был отрицательным (−$383 млн). Вывод: прибыль UAL подкреплена деньгами, но FCF после капзатрат скромный ({pc(V['fcf_m']['2025'])} выручки в 2025) и чувствителен к росту.")
doc.add_heading("Баланс и долг", 2)
IMG(charts["lev"], 15); SRC("8-K UAL; лист Metrics / Debt.")
P(f"[Ф] На 30.06.2026: долг, финансовая аренда и прочие финансовые обязательства — ${n1(V['debt']['1П26']/1000)} млрд, операционная аренда — ${n1(V['oplease']['1П26']/1000)} млрд, пенсии — ${n1(V['pension']['1П26']/1000)} млрд; итого скорр. долг ${n1(v('adj_debt','1П26')/1000)} млрд. Денежные средства и инвестиции — ${n1(v('cash_all','1П26')/1000)} млрд, чистый скорр. долг — ${n1(ND/1000)} млрд (совпадает с 8-K). Чистый долг/EBITDAR = {n1(v('nd_ebitdar','TTM 2К26'))}x по постоянному определению (2,2x по новому определению UAL). Покрытие процентов скорр. EBIT — {n1(V['cov']['TTM 2К26'])}x. Ликвидность $19,6 млрд (31% выручки), ковенант минимальной ликвидности — $2 млрд. Рейтинги BB+ (S&P, позитивный прогноз) / Ba1 / BB+. Платежи по долгу, аренде и прочим финансовым обязательствам (вкл. проценты) — ${n1(V['mat_tot']['2026'])} млрд в 2026, ${n1(V['mat_tot']['2027'])} млрд в 2027, ${n1(V['mat_tot']['2028'])} млрд в 2028 (10-K 2025). Выкуп: $637 млн в 2025, остаток программы $755 млн; во 2К26 выкупа не было.")
P(f"EBITDA или CFO? Долг/CFO за TTM ({n1(V['nd_cfo']['TTM 2К26'])}x) мягче, чем долг/EBITDAR ({n1(v('nd_ebitdar','TTM 2К26'))}x), потому что CFO включает приток авансов и почти нулевой денежный налог. Оба эффекта временные: авансы разворачиваются в спад, NOL заканчиваются. Для авиакомпании надёжнее чистый долг (с арендой и пенсиями) к EBITDAR, оценённому через цикл, с проверкой по CFO без изменения авансов ({n1(V['nd_cfoex']['2025'])}x в 2025).")
doc.add_heading("Чувствительность прибыли", 2)
P(f"[О] UAL не хеджирует топливо. ±10% к цене топлива = ∓${n1(-V['s4']['fuel_ttm']/1000)} млрд опер. прибыли за TTM ({pc(-V['s4']['fuel_pct_ttm'],0)} скорр. опер. прибыли ${n1(V['s4']['adj_ttm']/1000)} млрд) и ∓${n1(-V['s4']['fuel_27']/1000)} млрд в 2027E. ±1% RASM = ±${n2(V['s4']['rasm_ttm']/1000)} млрд ({pc(V['s4']['rasm_pct_ttm'],0)}). Это статический эффект: до участия в прибыли и без переноса в тариф (лист Sensitivity).")
T(sens, [4.2, 2.3, 2.3, 2.3, 2.3, 2.3]); SRC(sens_src)

doc.add_heading("3. Сравнительный анализ и оценка", 1)
doc.add_heading("Пиры и мультипликаторы", 2)
P("Почему такие пиры: Delta и American — сетевые перевозчики с хабами, премиальными кабинами и международными сетями, то есть прямые аналоги. Southwest, Alaska и JetBlue — внутренние и лоукост-модели, нужны для контекста сектора. Где сравнение некорректно: LUV летает на короткие рейсы без премиум-кабин; ALK интегрирует Hawaiian (убыток), JBLU убыточна; у AAL отрицательный капитал; у Delta свой НПЗ (дешевле топливо, выручка НПЗ в отчётности); TTM-прибыль UAL включает $615 млн разовых доходов от sale-leaseback. Air France-KLM, IAG и Lufthansa отчитываются по МСФО (IFRS 16: вся аренда уже в EBITDA), их мультипликаторы не рассчитывались — первичные данные за 1П26 не собраны.")
T(mult, [3.6, 1.7, 1.7, 1.7, 1.7, 1.7, 1.7, 1.8], 8); SRC(mult_src)
hist = [["", "2019", "2022", "2023", "2024", "2025", "Текущий"]]
for k, lab, f in [("hm_evebitdar", "EV/EBITDAR", lambda x: n1(x) + "x"), ("hm_pe", "P/E (скорр. EPS)", lambda x: n1(x) + "x"), ("hm_pb", "P/B", lambda x: n2(x) + "x"), ("hm_fcfy", "FCF yield", pc), ("hm_evrev", "EV/выручка", lambda x: n2(x) + "x")]:
    hist.append([lab] + [f(V[k][c]) for c in ["2019", "2022", "2023", "2024", "2025", "Текущий"]])
T(hist, [3.6, 2.0, 2.0, 2.0, 2.0, 2.0, 2.2]); SRC("лист Multiples; цены на конец года (Yahoo Finance), EBITDAR = скорр. EBITDA + аренда самолётов; текущий — цена 13.10.2026 и TTM 2К26.")
P(f"Против собственной истории UAL не дешёвая: EV/EBITDAR {n1(V['hm_evebitdar']['Текущий'])}x — у верхней границы диапазона {n1(v('hist_min','EV/EBITDAR'))}–{n1(v('hist_max','EV/EBITDAR'))}x. Но TTM EBITDAR занижен топливом. На нормализованной прибыли (средняя маржа цикла {pc(MID)} × выручка 2027E) EPS = ${n2(v('im_normeps'))}, и P/E = {n1(v('im_pe_norm'))}x против медианы истории {n1(v('hist_med','P/E'))}x — это дешевле. Оценка по медиане истории EV/EBITDAR × EBITDAR 2027E — ${v('im_hist27'):.0f}, по нормализованному P/E — ${v('im_normpe'):.0f}.")
doc.add_heading("Нормализация", 2)
P(f"[О] Средняя скорр. опер. маржа через цикл (2019, 2022–2025, без пандемийных 2020–2021) — {pc(MID)}; без слабого 2022 — {pc(v('midcycle_ex22','2025'))}. В базовом сценарии маржа восстанавливается с {pc(V['Base_m']['2026E'])} (2026E) до {pc(V['Base_m']['2032E'])} (2032E), то есть до средней за цикл, а не до пика. Прогноз не лучше истории.")
doc.add_heading("DCF (FCFF)", 2)
YF_ = ["2026E", "2027E", "2028E", "2029E", "2030E", "2031E", "2032E"]
dc = [["Допущение (база)", "2026E", "2027E", "2028E", "2029E", "2030E", "2031E", "2032E"]]
for k, lab, f in [("Base_in_asm", "Рост ASM", pc), ("Base_in_core", "Базовый рост TRASM", pc), ("Base_in_pt", "Перенос топлива", lambda x: pc(x, 0)), ("Base_in_casm", "Рост CASM-ex", pc), ("Base_in_fuel", "Топливо, $/галл.", n2)]:
    dc.append([lab] + [f(V[k][c]) for c in YF_])
dc.append(["Капзатраты, % выручки"] + [pc(V["Base_capex"]["2026E"] / V["Base_rev"]["2026E"])] + [pc(V["Base_in_capex"][c]) for c in YF_[1:]])
for k, lab, f in [("Base_rev", "Выручка, $ млрд", lambda x: n1(x / 1000)), ("Base_trasm", "TRASM, центы", n2), ("Base_m", "Скорр. опер. маржа", pc), ("Base_eps", "EPS скорр., $", n2), ("Base_fcff", "FCFF, $ млрд", lambda x: n2(x / 1000))]:
    dc.append([lab] + [f(V[k][c]) for c in YF_])
T(dc, [3.8] + [1.8] * 7, 8); SRC("листы Assumptions и DCF; допущения [Д], расчёт [О].")
P(f"Механика: выручка = ASM × TRASM. TRASM(t) = TRASM(t−1) × (1 + базовый рост) + доля переноса × изменение топлива на ASM — так рост цены топлива связан с тарифом. Расходы = CASM-ex × ASM + топливо (галлоны/ASM × цена) + сторонние услуги + участие в прибыли (12% EBIT). FCFF = NOPAT + D&A − скорр. капзатраты + приток оборотного капитала (10% прироста выручки). Аренда: операционная аренда включена в долг и EV, а её процентная часть (6,5% × обязательство) возвращена в NOPAT, чтобы не учесть её дважды. Налог — 22%; щит NOL ($2,3 млрд в налоговом выражении) добавлен отдельно как $2,0 млрд PV. Дата оценки — 30.06.2026, середина периода; из FCFF 2026E вычтен фактический FCFF 1П26 (${n1(v('h1fcff')/1000)} млрд).")
P(f"WACC = {pc(WACC)}: Rf 5,24% (UST10 на 09.10.2026), β {n2(BETA)} (60 мес., против S&P 500; по Блюму — {n2(V['beta_adj']['v'])}), ERP 5,0% [Д] → Ke {pc(KE)}. Kd 6,5% до налога [Д: UST10 + ~1,25 п.п. спред BB+; ноты 2031 размещены под 5,375% при UST10 4,29%] → 5,07% после налога. Веса по рынку: капитал ${n1(v('w_E')/1000)} млрд, скорр. долг с арендой и пенсиями ${n1(v('w_D')/1000)} млрд ({pc(v('w_wd'),0)}).")
P(f"Терминальная стоимость двумя способами. (1) Гордон: g = 2,5%, RONIC 10% → TV = NOPAT × (1 + g) × (1 − g/RONIC)/(WACC − g); подразумеваемый EV/EBITDAR = {n1(v('Base_impl_mult','2032E'))}x. (2) Выходной мультипликатор: медиана истории UAL {n1(v('hist_med','EV/EBITDAR'))}x (в пессимистичном — минимум {n1(v('hist_min','EV/EBITDAR'))}x, в оптимистичном — максимум {n1(v('hist_max','EV/EBITDAR'))}x). Результаты базы: ${n2(v('Base_psg','2032E'))} (Гордон) и ${n2(v('Base_psm','2032E'))} (мультипликатор), среднее ${n2(SV['Base'])}. Доля TV в EV — {pc(v('Base_tvshare','2032E'),0)}: FCFF отрицателен до 2029 из-за капзатрат.")
doc.add_heading("Reverse DCF", 2)
P(f"[О] При цене ${n2(PRICE)} модель подразумевает поправку к базовому росту TRASM {spc(v('rv_impl_d'),2)} в год, то есть операционную маржу 2032E {pc(RVM)} (средняя за цикл — {pc(MID)}), и WACC {pc(IW)} (наш — {pc(WACC)}). Рынок закладывает ровно возврат к средней марже цикла — реалистично, но без запаса. Каждые +0,5 п.п. роста TRASM в год сверх базы добавляют ~$50 на акцию (лист ReverseDCF).")
doc.add_heading("Сценарии", 2)
scn = [["", "Пессимистичный", "Базовый", "Оптимистичный", "Взвешенно"],
       ["Описание", "Рецессия 2027 (базовый TRASM −5%, ASM −2%), топливо $4,10", "Нефть по EIA ($3,30 в 2027), маржа к средней за цикл", "Премиальный спрос, топливо $2,90, маржа ~11–12%", ""],
       ["Вероятность", "25%", "50%", "25%", "100%"],
       ["EPS 2027E", usd(V["sc_eps27"]["Bear"], 2), usd(V["sc_eps27"]["Base"], 2), usd(V["sc_eps27"]["Bull"], 2), ""],
       ["Маржа 2032E", pc(V["sc_m32"]["Bear"]), pc(V["sc_m32"]["Base"]), pc(V["sc_m32"]["Bull"]), ""],
       ["Стоимость акции", usd(SV["Bear"], 2), usd(SV["Base"], 2), usd(SV["Bull"], 2), usd(SV["Взвешенно"], 2)],
       ["Потенциал", spc(V["sc_up"]["Bear"], 0), spc(V["sc_up"]["Base"], 0), spc(V["sc_up"]["Bull"], 0), spc(V["sc_up"]["Взвешенно"], 0)]]
T(scn, [3.0, 3.6, 3.6, 3.6, 2.4], 8.5); SRC("лист Scenarios.")
IMG(charts["ff"], 16); SRC("UAL_model.xlsx.")
P(f"Целевая цена = взвешенная DCF, округлённая до ${TGT:.0f}. Диапазон ${LOW:.0f}–{HIGH:.0f} = минимум и максимум из трёх методов (DCF база, медиана истории EV/EBITDAR × EBITDAR 2027E, нормализованный P/E). Цель не переносится на 12 месяцев вперёд по Ke — это консервативно: с переносом было бы ${v('t_roll'):.0f}.")
T(wg, [3, 2.4, 2.4, 2.4, 2.4, 2.4]); SRC("лист Sensitivity: стоимость акции, метод Гордона, базовый сценарий.")

doc.add_heading("4. Риски и катализаторы", 1)
T(risks, [5.4, 1.8, 1.8, 7.0], 8.5)
P("Катализаторы и сроки: отчёт за 3К26 и прогноз на 4К (20–21.10.2026) — проверка переноса топлива и спроса; движение нефти (EIA прогнозирует Brent $105 в 4К26 и $84 в 2027); инвестиционный рейтинг (цель — 2026); A321XLR на международных рейсах (начало 2027); Starlink во всём парке (конец 2027); возобновление выкупа; сертификация MAX 10 (2027).")
P("Что сделает рекомендацию неверной:", bold=True)
B("Апсайд (надо было покупать): топливо устойчиво ниже $3,00/галл. при сохранении TRASM г/г выше +5% в 1П27 и подтверждение UAL маржи до налога ≥10% на 2027 — это путь оптимистичного сценария (~$200).")
B("Даунсайд (надо было продавать): Brent выше $110 до весны 2027, TRASM 4К26 ниже +10% г/г, CASM-ex 2027 выше +4%, чистый долг/EBITDAR выше 3,0x — путь к пессимистичному сценарию (~$35).")

doc.add_heading("5. Рекомендация", 1)
P(f"{RATING}, целевая цена ${TGT:.0f} (диапазон ${LOW:.0f}–{HIGH:.0f}), потенциал {spc(UP)} к цене ${n2(PRICE)} на 13.10.2026, горизонт 12 месяцев.", bold=True)
P("Что отслеживать после рекомендации:", bold=True)
T(monitor, [5.5, 4, 6.5])
doc.add_heading("Ограничения анализа", 2)
T(limits, [7, 9])
P("Дополнительно: бета 1,36 — сырая; по Блюму 1,24 WACC снижается до ~8,2%, а оценка растёт на ~$10. Консенсус EPS взят с агрегатора (Yahoo Finance) и может быть устаревшим. Скорр. долг за 2019 — оценка по XBRL (пенсии в узком определении). Данные о долях хабов — вторичный источник (BTS через Simple Flying).", size=9)
doc.add_heading("Источники", 2)
for s_ in ["SEC EDGAR: UAL 10-K 2025 (12.02.2026), 10-Q 2К26 (16.07.2026), 8-K с пресс-релизами и Investor Updates 2020–2026; XBRL companyfacts UAL, DAL, AAL, LUV, ALK, JBLU.",
           "8-K за 2К26: Delta (10.07.2026), American (23.07.2026), Southwest (22.07.2026); Delta — прогноз 09.10.2026 (Reuters/CNBC).",
           "FRED: DGS10, DJFUELUSGULF, MJFUELUSGULF, DCOILBRENTEU; EIA STEO, октябрь 2026.",
           "Yahoo Finance chart API: цены на закрытие 13.10.2026 и месячные ряды для беты; консенсус EPS — Yahoo Finance Analysis.",
           "UAL IR: дата отчёта за 3К26 (20.10.2026); Morgan Stanley Laguna Conference 16.09.2026 (краткое изложение Investing.com).",
           "Simple Flying (BTS 2024) — доли хабов; The Points Guy — продление договора с Chase до 2029; PIRG «Plane Truth 2026» — пунктуальность."]:
    B(s_)
doc.save(f"{OUTDIR}/UAL_report.docx")
print("docx ok")
